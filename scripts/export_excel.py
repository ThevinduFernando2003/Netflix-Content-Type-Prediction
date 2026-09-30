"""Recover a string CSV from the Excel extract.

Excel stored a handful of titles as numbers, times, or dates (``3%``, ``23:59``,
``9-Feb``). This exporter writes the displayed text, not the coerced Python type.
"""

from __future__ import annotations

import csv
import shutil
import sys
from datetime import datetime, time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook
from src.config import repo_root
from src.logging_conf import get_logger

logger = get_logger(__name__)

EXCEL_NAME = "Dataset.csv.xlsx"
CSV_RELATIVE = Path("data/raw/netflix_titles.csv")


def _format_cell(value: object, number_format: str | None, column: str) -> str:
    if value is None:
        return ""
    fmt = number_format or "General"

    if isinstance(value, datetime):
        if "mmm" in fmt.lower():
            return f"{value.day}-{value.strftime('%b')}"
        if column == "date_added":
            return f"{value.month}/{value.day}/{value.year}"
        return f"{value.day}-{value.strftime('%b')}"

    if isinstance(value, time):
        return f"{value.hour}:{value.minute:02d}"

    if isinstance(value, float):
        if "%" in fmt:
            percent = value * 100
            if percent == int(percent):
                return f"{int(percent)}%"
            return f"{percent:g}%"
        if value == int(value):
            as_int = int(value)
            if "#,##0" in fmt:
                return f"{as_int:,}"
            return str(as_int)
        return str(value)

    if isinstance(value, int):
        if "#,##0" in fmt:
            return f"{value:,}"
        return str(value)

    return str(value).strip()


def export_excel_to_csv(excel_path: Path, csv_path: Path) -> int:
    """Write ``csv_path`` from ``excel_path``. Returns the number of data rows."""
    workbook = load_workbook(excel_path, data_only=True, read_only=False)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows()
    header_cells = next(rows)
    header = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        for row in rows:
            if all(cell.value is None for cell in row):
                continue
            values = [
                _format_cell(cell.value, cell.number_format, column)
                for cell, column in zip(row, header, strict=True)
            ]
            writer.writerow(values)
            count += 1
    workbook.close()
    logger.info("Wrote %s rows to %s", count, csv_path)
    return count


def main() -> None:
    root = repo_root()
    candidates = [
        root / EXCEL_NAME,
        root / "data" / "raw" / EXCEL_NAME,
    ]
    excel_path = next((path for path in candidates if path.exists()), None)
    if excel_path is None:
        raise FileNotFoundError("Dataset.csv.xlsx was not found at the repo root or in data/raw/.")
    destination = root / "data" / "raw" / EXCEL_NAME
    destination.parent.mkdir(parents=True, exist_ok=True)
    if excel_path.resolve() != destination.resolve():
        shutil.copy2(excel_path, destination)
        logger.info("Copied Excel source to %s", destination)
    export_excel_to_csv(destination, root / CSV_RELATIVE)


if __name__ == "__main__":
    main()
